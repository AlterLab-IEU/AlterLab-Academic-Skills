#!/usr/bin/env python3
"""build_deck.py - build an editable 16:9 PowerPoint deck from a JSON plan (keyless).

Claude writes the plan (titles, bullets, figure paths, citations, speaker notes) and this
script lays it out with python-pptx in one consistent theme: no image model, no API key, and
every word stays editable. Export to PDF with LibreOffice and review the pages as images:

    uv run --python ">=3.11" --with python-pptx --with pillow python scripts/build_deck.py deck.json -o talk.pptx
    soffice --headless --convert-to pdf --outdir . talk.pptx
    uv run --python ">=3.11" --with pymupdf python scripts/pdf_to_images.py talk.pdf review/slide

Plan format (paths are relative to the plan file):

    {"theme": {"primary": "#1B2A4A", "accent": "#2A9D8F", "text": "#222222",
               "background": "#FFFFFF", "font": "Calibri"},
     "footer": "A. Researcher - Conference 2026",
     "slides": [
       {"layout": "title", "title": "...", "subtitle": "Name, Affiliation", "notes": "..."},
       {"layout": "section", "title": "Methods"},
       {"layout": "bullets", "title": "...", "bullets": ["...", "..."], "image": "figures/a.png",
        "citation": "Smith et al., 2024", "notes": "..."},
       {"layout": "image", "title": "...", "image": "figures/result.png", "caption": "...",
        "citation": "..."},
       {"layout": "two_column", "title": "...", "left": ["..."], "right": ["..."],
        "left_heading": "...", "right_heading": "..."},
       {"layout": "takeaway", "title": "...", "message": "One sentence the audience keeps"}]}

Checks (printed as warnings, never silently fixed): more than 4 bullets, a bullet over 12 words,
a missing image, a title over 60 characters, a body font under 18 pt. A missing image is an
error unless --allow-missing-images.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt
except ImportError:  # pragma: no cover - reported to the user
    sys.exit("build_deck.py needs python-pptx: `uv run --with python-pptx --with pillow python ...`")

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
DEFAULT_THEME = {"primary": "#1B2A4A", "accent": "#2A9D8F", "text": "#222222",
                 "background": "#FFFFFF", "muted": "#666666", "font": "Calibri",
                 "title_pt": 36, "body_pt": 24, "small_pt": 14}


def _rgb(hex_colour: str) -> RGBColor:
    return RGBColor.from_string(hex_colour.lstrip("#").upper()[:6])


class Deck:
    def __init__(self, theme: dict, footer: str, base: Path, allow_missing: bool):
        self.t = {**DEFAULT_THEME, **(theme or {})}
        self.footer = footer
        self.base = base
        self.allow_missing = allow_missing
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.warnings: list[str] = []
        self.n = 0
        if self.t["body_pt"] < 18:
            self.warnings.append("theme body_pt is under 18 pt: unreadable from the back of a room")

    # ----------------------------------------------------------------- primitives
    def _text(self, slide, left, top, width, height, text, size, colour=None, bold=False,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            run = p.add_run()
            run.text = line
            f = run.font
            f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, self.t["font"]
            f.color.rgb = _rgb(colour or self.t["text"])
        return box

    def _bullets(self, slide, left, top, width, height, items):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(14)
            run = p.add_run()
            run.text = "•  " + item
            f = run.font
            f.size, f.name = Pt(self.t["body_pt"]), self.t["font"]
            f.color.rgb = _rgb(self.t["text"])

    def _image(self, slide, path_text, left, top, max_w, max_h):
        path = (self.base / path_text).resolve()
        if not path.is_file():
            msg = f"slide {self.n}: image not found: {path_text}"
            if not self.allow_missing:
                raise FileNotFoundError(msg)
            self.warnings.append(msg)
            self._text(slide, left, top, max_w, max_h, f"[missing image: {path_text}]",
                       self.t["small_pt"], colour=self.t["muted"], align=PP_ALIGN.CENTER,
                       anchor=MSO_ANCHOR.MIDDLE)
            return
        from PIL import Image

        with Image.open(path) as im:
            iw, ih = im.size
        scale = min(max_w / iw, max_h / ih)
        w, h = int(iw * scale), int(ih * scale)
        slide.shapes.add_picture(str(path), left + (max_w - w) // 2, top + (max_h - h) // 2,
                                 Emu(w), Emu(h))

    def _frame(self, spec, dark=False):
        self.n += 1
        slide = self.prs.slides.add_slide(self.blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(self.t["primary"] if dark else self.t["background"])
        title = spec.get("title", "")
        if len(title) > 60:
            self.warnings.append(f"slide {self.n}: title is {len(title)} characters (aim for <= 60)")
        if spec.get("notes"):
            slide.notes_slide.notes_text_frame.text = spec["notes"]
        if not dark:
            if self.footer:
                self._text(slide, MARGIN, H - Inches(0.5), Inches(9), Inches(0.35), self.footer,
                           self.t["small_pt"], colour=self.t["muted"])
            self._text(slide, W - Inches(1.4), H - Inches(0.5), Inches(0.8), Inches(0.35),
                       str(self.n), self.t["small_pt"], colour=self.t["muted"], align=PP_ALIGN.RIGHT)
        if spec.get("citation"):
            self._text(slide, MARGIN, H - Inches(0.9), W - 2 * MARGIN, Inches(0.35),
                       spec["citation"], self.t["small_pt"], colour=self.t["muted"], italic=True)
        return slide

    def _title_bar(self, slide, title):
        self._text(slide, MARGIN, Inches(0.35), W - 2 * MARGIN, Inches(0.9), title,
                   self.t["title_pt"], colour=self.t["primary"], bold=True, anchor=MSO_ANCHOR.MIDDLE)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.3), Inches(1.2), Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(self.t["accent"])
        bar.line.fill.background()

    def _check_bullets(self, items):
        if len(items) > 4:
            self.warnings.append(f"slide {self.n}: {len(items)} bullets (aim for 3-4)")
        for b in items:
            if len(b.split()) > 12:
                self.warnings.append(f"slide {self.n}: bullet over 12 words: {b[:50]!r}")

    # ----------------------------------------------------------------- layouts
    def add(self, spec: dict):
        layout = spec.get("layout", "bullets")
        top = Inches(1.6)
        body_h = H - top - Inches(1.1)
        if layout == "title":
            s = self._frame(spec, dark=True)
            self._text(s, MARGIN, Inches(2.2), W - 2 * MARGIN, Inches(1.8), spec.get("title", ""),
                       44, colour="#FFFFFF", bold=True, anchor=MSO_ANCHOR.BOTTOM)
            self._text(s, MARGIN, Inches(4.2), W - 2 * MARGIN, Inches(1.4),
                       spec.get("subtitle", ""), 24, colour=self.t["accent"])
        elif layout == "section":
            s = self._frame(spec, dark=True)
            self._text(s, MARGIN, Inches(2.8), W - 2 * MARGIN, Inches(1.5), spec.get("title", ""),
                       40, colour="#FFFFFF", bold=True, anchor=MSO_ANCHOR.MIDDLE)
        elif layout == "takeaway":
            s = self._frame(spec)
            self._title_bar(s, spec.get("title", "Take-home message"))
            self._text(s, Inches(1.2), Inches(2.4), W - Inches(2.4), Inches(2.8),
                       spec.get("message", ""), 32, colour=self.t["primary"], bold=True,
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        elif layout == "image":
            s = self._frame(spec)
            self._title_bar(s, spec.get("title", ""))
            cap_h = Inches(0.6) if spec.get("caption") else 0
            self._image(s, spec["image"], MARGIN, top, W - 2 * MARGIN, body_h - cap_h)
            if spec.get("caption"):
                self._text(s, MARGIN, top + body_h - cap_h, W - 2 * MARGIN, cap_h,
                           spec["caption"], 18, colour=self.t["text"], align=PP_ALIGN.CENTER)
        elif layout == "two_column":
            s = self._frame(spec)
            self._title_bar(s, spec.get("title", ""))
            col_w = (W - 2 * MARGIN - Inches(0.5)) // 2
            for i, side in enumerate(("left", "right")):
                x = MARGIN + i * (col_w + Inches(0.5))
                y = top
                if spec.get(f"{side}_heading"):
                    self._text(s, x, y, col_w, Inches(0.6), spec[f"{side}_heading"], 26,
                               colour=self.t["accent"], bold=True)
                    y += Inches(0.7)
                items = spec.get(side, [])
                self._check_bullets(items)
                self._bullets(s, x, y, col_w, body_h - (y - top), items)
        else:  # bullets, optionally with an image on the right
            s = self._frame(spec)
            self._title_bar(s, spec.get("title", ""))
            items = spec.get("bullets", [])
            self._check_bullets(items)
            if spec.get("image"):
                text_w = Inches(5.6)
                self._bullets(s, MARGIN, top, text_w, body_h, items)
                self._image(s, spec["image"], MARGIN + text_w + Inches(0.3), top,
                            W - 2 * MARGIN - text_w - Inches(0.3), body_h)
            else:
                self._bullets(s, MARGIN, top, W - 2 * MARGIN, body_h, items)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("plan", type=Path, help="deck plan (JSON)")
    ap.add_argument("-o", "--output", type=Path, required=True, help="output .pptx")
    ap.add_argument("--allow-missing-images", action="store_true",
                    help="put a placeholder instead of failing when a figure file is missing")
    args = ap.parse_args(argv)
    try:
        plan = json.loads(args.plan.read_text(encoding="utf-8"))
        deck = Deck(plan.get("theme"), plan.get("footer", ""), args.plan.parent,
                    args.allow_missing_images)
        for spec in plan["slides"]:
            deck.add(spec)
    except (OSError, ValueError, KeyError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    args.output.parent.mkdir(parents=True, exist_ok=True)
    deck.prs.save(args.output)
    print(json.dumps({"output": str(args.output), "slides": deck.n, "warnings": deck.warnings},
                     indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
