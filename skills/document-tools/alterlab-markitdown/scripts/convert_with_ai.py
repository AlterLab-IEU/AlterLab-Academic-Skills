#!/usr/bin/env python3
"""
Convert documents to Markdown with descriptions of the images inside them.

Two backends:

* ``claude`` (keyless; the default when no OPENROUTER_API_KEY is set). MarkItDown
  converts the text, and every embedded image (PPTX slides, DOCX/XLSX media, PDF
  images when PyMuPDF is installed, or the input image itself) is written to
  ``<output stem>_images/`` with a manifest ``<output stem>_images.json`` that
  carries the chosen description prompt. Claude then Reads each image and writes
  the descriptions into the Markdown, so no second model or API key is involved.
* ``openrouter`` (optional key). MarkItDown calls a vision model through OpenRouter
  (OpenAI-compatible client) and writes the descriptions inline.

``--backend auto`` uses ``openrouter`` when a key is available and ``claude``
otherwise, and says which on stderr. ``--backend openrouter`` without a key exits 2
instead of silently converting without descriptions.
"""

import argparse
import json
import os
import re
import sys
import zipfile
from pathlib import Path


# AlterLab model convention — default reviewed 2026-09-23; override via ALTERLAB_MODEL.
# See skills/core/shared/model_env.md before changing the default. ALTERLAB_MODEL holds an
# Anthropic model ID; openrouter_slug() derives the "provider/model" slug OpenRouter needs.
DEFAULT_MODEL = "claude-opus-5-5"


def alterlab_model() -> str:
    """Return the model ID to use: $ALTERLAB_MODEL if set/non-empty, else the dated default."""
    return os.environ.get("ALTERLAB_MODEL") or DEFAULT_MODEL


def openrouter_slug(model: str) -> str:
    """Map an Anthropic model ID to its OpenRouter slug; pass anything else through.

    OpenRouter names Claude models with a provider prefix and a dotted version
    (claude-opus-5-5 -> anthropic/claude-opus-5.5). IDs that already carry a provider
    prefix ("anthropic/...", "google/...") are returned unchanged.
    """
    if "/" in model:
        return model
    m = re.fullmatch(r"claude-([a-z]+)-(\d+)(?:-(\d{1,2}))?(?:-\d{8})?", model)
    if not m:
        return model
    family, major, minor = m.groups()
    return f"anthropic/claude-{family}-{major}" + (f".{minor}" if minor else "")


# Predefined prompts for different use cases
PROMPTS = {
    'scientific': """
Analyze this scientific image or diagram. Provide:
1. Type of visualization (graph, chart, microscopy, diagram, etc.)
2. Key data points, trends, or patterns
3. Axes labels, legends, and scales
4. Notable features or findings
5. Scientific context and significance
Be precise, technical, and detailed.
    """.strip(),
    
    'presentation': """
Describe this presentation slide image. Include:
1. Main visual elements and their arrangement
2. Key points or messages conveyed
3. Data or information presented
4. Visual hierarchy and emphasis
Keep the description clear and informative.
    """.strip(),
    
    'general': """
Describe this image in detail. Include:
1. Main subjects and objects
2. Visual composition and layout
3. Text content (if any)
4. Notable details
5. Overall context and purpose
Be comprehensive and accurate.
    """.strip(),
    
    'data_viz': """
Analyze this data visualization. Provide:
1. Type of chart/graph (bar, line, scatter, pie, etc.)
2. Variables and axes
3. Data ranges and scales
4. Key patterns, trends, or outliers
5. Statistical insights
Focus on quantitative accuracy.
    """.strip(),
    
    'medical': """
Describe this medical image. Include:
1. Type of medical imaging (X-ray, MRI, CT, microscopy, etc.)
2. Anatomical structures visible
3. Notable findings or abnormalities
4. Image quality and contrast
5. Clinical relevance
Be professional and precise.
    """.strip()
}


def convert_with_ai(
    input_file: Path,
    output_file: Path,
    api_key: str,
    model: str = None,
    prompt_type: str = "general",
    custom_prompt: str = None
) -> bool:
    """
    Convert a file to Markdown with AI image descriptions.
    
    Args:
        input_file: Path to input file
        output_file: Path to output Markdown file
        api_key: OpenRouter API key
        model: Model name (default: $ALTERLAB_MODEL or the dated DEFAULT_MODEL constant)
        prompt_type: Type of prompt to use
        custom_prompt: Custom prompt (overrides prompt_type)

    Returns:
        True if successful, False otherwise
    """
    if not model:
        model = alterlab_model()
    model = openrouter_slug(model)
    try:
        from markitdown import MarkItDown
        from openai import OpenAI

        # Initialize OpenRouter client (OpenAI-compatible)
        client = OpenAI(
            api_key=api_key,
            base_url="https://openrouter.ai/api/v1"
        )
        
        # Select prompt
        if custom_prompt:
            prompt = custom_prompt
        else:
            prompt = PROMPTS.get(prompt_type, PROMPTS['general'])
        
        print(f"Using model: {model}")
        print(f"Prompt type: {prompt_type if not custom_prompt else 'custom'}")
        print(f"Converting: {input_file}")
        
        # Create MarkItDown with AI support
        md = MarkItDown(
            llm_client=client,
            llm_model=model,
            llm_prompt=prompt
        )
        
        # Convert file
        result = md.convert(str(input_file))
        
        # Create output with metadata
        content = f"# {result.title or input_file.stem}\n\n"
        content += f"**Source**: {input_file.name}\n"
        content += f"**Format**: {input_file.suffix}\n"
        content += f"**AI Model**: {model}\n"
        content += f"**Prompt Type**: {prompt_type if not custom_prompt else 'custom'}\n\n"
        content += "---\n\n"
        content += result.text_content
        
        # Write output
        output_file.parent.mkdir(parents=True, exist_ok=True)
        output_file.write_text(content, encoding='utf-8')
        
        print(f"✓ Successfully converted to: {output_file}")
        return True
        
    except Exception as e:
        print(f"✗ Error: {str(e)}", file=sys.stderr)
        return False


IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}


def extract_images(input_file: Path, image_dir: Path) -> list:
    """Write the images embedded in input_file to image_dir; return manifest rows.

    PPTX pictures are mapped to their slide with python-pptx (installed by
    markitdown[pptx]); otherwise DOCX/XLSX/PPTX media are read from the zip's media
    folder. PDF images need PyMuPDF. An image input is copied as-is.
    """
    suffix = input_file.suffix.lower()
    rows = []

    def _save(name, blob, **where):
        image_dir.mkdir(parents=True, exist_ok=True)
        out = image_dir / name
        out.write_bytes(blob)
        rows.append({"image": str(out), "bytes": len(blob), **where})

    if suffix in IMAGE_SUFFIXES:
        _save(f"image_01{suffix}", input_file.read_bytes(), source="input file")
        return rows
    if suffix == ".pptx":
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError:
            Presentation = None
        if Presentation is not None:
            def _pictures(shapes):
                for shape in shapes:
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        yield from _pictures(shape.shapes)
                    elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        yield shape

            for s_no, slide in enumerate(Presentation(str(input_file)).slides, 1):
                title = slide.shapes.title.text if slide.shapes.title is not None else ""
                for i, pic in enumerate(_pictures(slide.shapes), 1):
                    _save(f"slide{s_no:02d}_img{i:02d}.{pic.image.ext}", pic.image.blob,
                          slide=s_no, slide_title=title, shape_name=pic.name)
            return rows
    if suffix in (".pptx", ".docx", ".xlsx"):
        media = {".pptx": "ppt/media/", ".docx": "word/media/", ".xlsx": "xl/media/"}[suffix]
        with zipfile.ZipFile(input_file) as z:
            names = sorted(n for n in z.namelist() if n.startswith(media))
            for i, name in enumerate(names, 1):
                _save(f"media{i:02d}{Path(name).suffix.lower()}", z.read(name), source=name)
        return rows
    if suffix == ".pdf":
        try:
            import pymupdf as fitz  # PyMuPDF >= 1.24
        except ImportError:
            try:
                import fitz  # older PyMuPDF
            except ImportError:
                print("note: PDF images are not extracted without PyMuPDF (uv run --with "
                      "pymupdf); Claude can Read the PDF pages directly instead.", file=sys.stderr)
                return rows
        doc = fitz.open(str(input_file))
        for p_no, page in enumerate(doc, 1):
            for i, img in enumerate(page.get_images(full=True), 1):
                info = doc.extract_image(img[0])
                if info and info.get("image"):
                    _save(f"page{p_no:02d}_img{i:02d}.{info.get('ext', 'png')}", info["image"],
                          page=p_no, width=info.get("width"), height=info.get("height"))
        if not rows:
            print("note: this PDF embeds no raster images; figures drawn as vectors are part of "
                  "the page - Read the PDF pages to describe them.", file=sys.stderr)
    return rows


def convert_keyless(input_file: Path, output_file: Path, prompt: str) -> bool:
    """Convert with MarkItDown (no LLM) and stage the images for Claude to describe."""
    from markitdown import MarkItDown

    result = MarkItDown().convert(str(input_file))
    image_dir = output_file.with_name(f"{output_file.stem}_images")
    rows = extract_images(input_file, image_dir)
    content = f"# {result.title or input_file.stem}\n\n"
    content += f"**Source**: {input_file.name}\n"
    content += f"**Format**: {input_file.suffix}\n\n"
    content += "---\n\n"
    content += result.markdown
    output_file.parent.mkdir(parents=True, exist_ok=True)
    output_file.write_text(content, encoding="utf-8")
    manifest = output_file.with_name(f"{output_file.stem}_images.json")
    manifest.write_text(json.dumps({"markdown": str(output_file), "prompt": prompt,
                                    "images": rows}, indent=2), encoding="utf-8")
    print(f"✓ Converted to: {output_file}")
    print(f"✓ {len(rows)} image(s) staged in {image_dir} (manifest: {manifest})")
    if rows:
        print("Next: Read each image in the manifest, describe it with the manifest's prompt, "
              "and add each description to the Markdown next to its slide/page.", file=sys.stderr)
    return True


def _key_from_env_file():
    """OPENROUTER_API_KEY from a .env file in the current or a parent directory."""
    for parent in [Path.cwd(), *Path.cwd().parents]:
        env = parent / ".env"
        if env.is_file():
            for line in env.read_text(encoding="utf-8", errors="replace").splitlines():
                if line.startswith("OPENROUTER_API_KEY="):
                    value = line.split("=", 1)[1].strip().strip("'\"")
                    if value:
                        return value
    return None


def main():
    parser = argparse.ArgumentParser(
        description="Convert documents to Markdown with AI-enhanced image descriptions",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=f"""
Available prompt types:
  scientific    - For scientific diagrams, graphs, and charts
  presentation  - For presentation slides
  general       - General-purpose image description
  data_viz      - For data visualizations and charts
  medical       - For medical imaging

Backends:
  auto (default)  openrouter when OPENROUTER_API_KEY is set, else claude
  claude          keyless: convert, then stage the images for Claude to describe with Read
  openrouter      MarkItDown calls a vision model through OpenRouter (needs the key)

Examples:
  # Keyless: convert slides and stage their images for Claude to describe
  python convert_with_ai.py slides.pptx slides.md --backend claude --prompt-type data_viz

  # Convert a scientific paper
  python convert_with_ai.py paper.pdf output.md --prompt-type scientific
  
  # Convert a presentation with custom model
  python convert_with_ai.py slides.pptx slides.md --model anthropic/claude-opus-5.5 --prompt-type presentation

  # Use custom prompt with advanced vision model
  python convert_with_ai.py diagram.png diagram.md --model anthropic/claude-opus-5.5 --custom-prompt "Describe this technical diagram"

  # Set API key via environment variable
  export OPENROUTER_API_KEY="sk-or-v1-..."
  python convert_with_ai.py image.jpg image.md

Environment Variables:
  OPENROUTER_API_KEY    Optional OpenRouter API key for --backend openrouter (or --api-key)
  ALTERLAB_MODEL        Override the default model ID (see skills/core/shared/model_env.md)

Popular Models (use with --model):
  anthropic/claude-opus-5.5 - Recommended for scientific vision (the dated default,
                              claude-opus-5-5, maps to this slug)
  Any OpenRouter vision slug (e.g. a current Google Gemini Pro Vision model) also works.
        """
    )
    
    parser.add_argument('input', type=Path, help='Input file')
    parser.add_argument('output', type=Path, help='Output Markdown file')
    parser.add_argument(
        '--backend', '-b',
        choices=['auto', 'claude', 'openrouter'],
        default='auto',
        help='auto (default): openrouter when a key is set, else the keyless claude backend'
    )
    parser.add_argument(
        '--api-key', '-k',
        help='OpenRouter API key for --backend openrouter (or set OPENROUTER_API_KEY env var)'
    )
    parser.add_argument(
        '--model', '-m',
        default=alterlab_model(),
        help='Model to use via OpenRouter (default: $ALTERLAB_MODEL or the dated DEFAULT_MODEL)'
    )
    parser.add_argument(
        '--prompt-type', '-t',
        choices=list(PROMPTS.keys()),
        default='general',
        help='Type of prompt to use (default: general)'
    )
    parser.add_argument(
        '--custom-prompt', '-p',
        help='Custom prompt (overrides --prompt-type)'
    )
    parser.add_argument(
        '--list-prompts', '-l',
        action='store_true',
        help='List available prompt types and exit'
    )
    
    args = parser.parse_args()
    
    # List prompts and exit
    if args.list_prompts:
        print("Available prompt types:\n")
        for name, prompt in PROMPTS.items():
            print(f"[{name}]")
            print(prompt)
            print("\n" + "="*60 + "\n")
        sys.exit(0)
    
    # Validate input file
    if not args.input.exists():
        print(f"Error: Input file '{args.input}' does not exist")
        sys.exit(1)

    api_key = args.api_key or os.environ.get('OPENROUTER_API_KEY') or _key_from_env_file()
    backend = args.backend
    if backend == 'openrouter' and not api_key:
        print("Error: --backend openrouter needs OPENROUTER_API_KEY (or --api-key); not falling "
              "back to a conversion without descriptions. Keyless alternative: --backend claude "
              "(stages the images for Claude to describe with Read).", file=sys.stderr)
        sys.exit(2)
    if backend == 'auto':
        backend = 'openrouter' if api_key else 'claude'
        if backend == 'claude':
            print("note: no OPENROUTER_API_KEY - keyless backend: converting with MarkItDown and "
                  "staging the images for Claude to describe.", file=sys.stderr)
    if backend == 'claude':
        prompt = args.custom_prompt or PROMPTS.get(args.prompt_type, PROMPTS['general'])
        try:
            success = convert_keyless(args.input, args.output, prompt)
        except Exception as e:  # noqa: BLE001 - report and exit non-zero
            print(f"✗ Error: {e}", file=sys.stderr)
            success = False
        sys.exit(0 if success else 1)

    # Convert file
    success = convert_with_ai(
        input_file=args.input,
        output_file=args.output,
        api_key=api_key,
        model=args.model,
        prompt_type=args.prompt_type,
        custom_prompt=args.custom_prompt
    )
    
    sys.exit(0 if success else 1)


if __name__ == '__main__':
    main()

